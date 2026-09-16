"""Build the Word overview using python-docx (available in the sibling RAG environment)."""

import json
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "reports"
INVENTORY = json.loads((ASSETS / "document_inventory.json").read_text())
DOCUMENT = ROOT / "Market_Lens_Project_Overview.docx"


def shade(cell, color):
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(element)


def para(doc, text, style=None):
    return doc.add_paragraph(text, style)


def bullets(doc, items):
    for item in items:
        para(doc, item, "List Bullet")


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Shading Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, value in zip(t.rows[0].cells, headers):
        cell.text = value
        shade(cell, "18344F")
        for r in cell.paragraphs[0].runs:
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.bold = True
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(5)
                for r in p.runs:
                    r.font.size = Pt(9)
        no_split = OxmlElement("w:cantSplit")
        t.rows[-1]._tr.get_or_add_trPr().append(no_split)
    para(doc, "")
    return t


def quote(doc, text):
    p = para(doc, text, "Quote")
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.12)
    for r in p.runs:
        r.font.size = Pt(9.5)
    return p


def code(doc, text):
    p = para(doc, text)
    for r in p.runs:
        r.font.name = "Consolas"
        r.font.size = Pt(9)
    p.paragraph_format.space_after = Pt(10)


def heading(doc, title, level=1):
    doc.add_heading(title, level)


def link(doc, label, url):
    p = doc.add_paragraph()
    hyperlink = OxmlElement("w:hyperlink")
    relationship = p.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink.set(qn("r:id"), relationship)
    run, props = OxmlElement("w:r"), OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2467AB")
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label + " — " + url
    run.append(text)
    hyperlink.append(run)
    p._p.append(hyperlink)


def build():
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.27), Inches(11.69)
    s.top_margin = s.bottom_margin = Inches(0.7)
    s.left_margin = s.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.color.rgb = RGBColor.from_string("18344F")
    doc.styles["Heading 1"].font.size = Pt(19)
    doc.styles["Heading 2"].font.size = Pt(13)
    doc.core_properties.title = "Market Lens — Project Overview and Development Journey"
    doc.core_properties.subject = "LangGraph multi-agent SaaS competitor research"
    doc.core_properties.author = "Agentic_AI_Project"
    header = s.header.paragraphs[0]
    header.text = "MARKET LENS  |  AGENTIC_AI_PROJECT"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string("647789")
    footer = s.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Project overview  •  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    para(doc, "PROJECT OVERVIEW & DEVELOPMENT JOURNEY", "Subtitle")
    doc.add_heading("Market Lens", 0)
    para(doc, "A multi-agent competitive intelligence system built with LangGraph", "Subtitle")
    para(doc, "SaaS-focused research • You.com search • OpenAI analysis • Streamlit interface")
    doc.add_picture(str(ASSETS / "architecture.png"), width=Inches(6.65))
    para(doc, "Prepared for Agentic_AI_Project")
    para(doc, "Prepared: " + datetime.now(timezone.utc).strftime("%d %B %Y, %H:%M UTC"))
    para(
        doc,
        "Scope: implemented architecture, application prompts, recorded development iterations, operational observations and learnings.",
    )
    para(
        doc,
        "Evidence basis: current source code, README, test outcomes recorded during development, our conversation, and local run metadata. No API keys or raw provider error bodies are included.",
    )
    doc.add_page_break()

    heading(doc, "1. Project overview")
    para(
        doc,
        "Market Lens turns a company name, domain and market scope into a structured competitor analysis briefing. It discovers and ranks competitors, gathers current web and news evidence, extracts comparable facts, independently reviews support, and presents findings with citations and visible research gaps.",
    )
    para(
        doc,
        "The first implementation targets business software and SaaS. Industry descriptions, discovery terms, feature taxonomies and pricing rules are registered separately, allowing the orchestration framework to support other industries later. New industry-specific categories would require coordinated schema, quality-gate and UI changes.",
    )
    heading(doc, "Intended users", 2)
    bullets(
        doc,
        [
            "Product teams: compare capabilities, packaging and potential differentiation.",
            "Strategy teams and founders: identify competing customer jobs, positioning and emerging threats.",
            "Consultants: produce repeatable, inspectable research snapshots with supporting evidence.",
        ],
    )
    heading(doc, "Inputs and outputs", 2)
    table(
        doc,
        ["Inputs", "Outputs"],
        [
            (
                "Company name and official domain",
                "Verified target baseline and ranked candidate list",
            ),
            (
                "Product/use case, geography and customer segment",
                "Three competitor profiles where sufficiently supported",
            ),
            (
                "News window and search/model/time budgets",
                "Pricing, features, positioning and recent news",
            ),
            (
                "Demo/live mode and optional human feedback",
                "Comparison table, strategic interpretations, Markdown/JSON exports and agent trace",
            ),
        ],
    )
    para(
        doc,
        "The offline demo uses fictional companies and facts. It exercises the real graph, budgets and review controls without paid calls. Live mode uses You.com and OpenAI; a completed run indicates workflow completion, not guaranteed factual correctness or exhaustive market coverage.",
    )

    heading(doc, "2. System architecture")
    para(
        doc,
        "The UI delegates execution to a research service rather than blocking the Streamlit render loop. The service runs the LangGraph pipeline in background workers, persists metadata and checkpoints, and exposes recorded events for UI polling. Typed schemas constrain the outputs agents exchange.",
    )
    table(
        doc,
        ["Component", "Responsibility", "Implementation"],
        [
            (
                "Streamlit UI",
                "Inputs, progress, review decisions, analysis and downloads",
                "app.py",
            ),
            (
                "Research service",
                "Background execution, saved-run resume and cancellation",
                "market_research/service.py",
            ),
            (
                "LangGraph pipeline",
                "Parent orchestration, conditional routing and competitor branches",
                "market_research/graph.py",
            ),
            (
                "Provider adapters",
                "Search, structured model calls, retries and concurrency",
                "market_research/providers.py",
            ),
            (
                "Quality engine",
                "Claim grounding, confidence, freshness and coverage",
                "market_research/quality.py",
            ),
            (
                "Persistent storage",
                "Atomic budgets, cache, events and run metadata",
                "market_research/storage.py + SQLite",
            ),
        ],
    )
    heading(doc, "Agent responsibilities", 2)
    table(
        doc,
        ["Agent", "Main work", "Structured result"],
        [
            (
                "Orchestrator",
                "Plan scope, dispatch branches, assess results and synthesize",
                "ResearchPlan / Briefing",
            ),
            ("Competitor Discovery", "Identify target and rank candidate overlap", "Discovery"),
            (
                "Research Collector",
                "Plan focused queries and gather sources",
                "CollectionPlan / Evidence records",
            ),
            (
                "Analysis Extractor",
                "Extract cited pricing, features, positioning and news",
                "Extraction / Claim",
            ),
            (
                "Evidence Reviewer",
                "Check selection identities and extracted claim support",
                "DiscoveryReview / Review",
            ),
        ],
    )
    para(
        doc,
        "These are application agents and graph nodes. No separate coding-assistant subagents were spawned during implementation. The application itself demonstrates delegation through parent-to-branch dispatch and evidence handoffs.",
    )
    doc.add_page_break()
    heading(doc, "3. LangGraph orchestration and feedback")
    doc.add_picture(str(ASSETS / "langgraph_workflow.png"), width=Inches(6.65))
    para(
        doc,
        "Figure 2. Parent graph, parallel competitor subgraphs, bounded follow-up and human review. Two human boxes show the selection and evidence contexts of the same parent human node.",
        "Caption",
    )
    para(
        doc,
        "Send launches up to three competitor branches. Each subgraph cycles through collect → extract → review, then either gathers targeted follow-up evidence or finishes. A reducer merges profiles by competitor domain, preventing one branch from overwriting another. The parent assesses the combined results before synthesis.",
    )
    para(
        doc,
        "The graph, rather than a free-form agent conversation, enforces mandatory categories, route choices, retry budgets and stop conditions. Model-generated plans can tailor queries but cannot silently remove a required category or increase limits.",
    )

    heading(doc, "4. APIs, libraries and configuration")
    table(
        doc,
        ["API / technology", "Use in this project", "Key implementation detail"],
        [
            (
                "You.com Search API",
                "Primary competitor discovery and web/news evidence",
                "POST https://ydc-index.io/v1/search; X-API-Key authentication",
            ),
            (
                "You.com extraction within Search",
                "Retrieve current Markdown page content",
                "full_page + extraction_source=fetch; snippets remain explicitly labeled on fallback",
            ),
            (
                "OpenAI Responses API",
                "Planning, discovery, extraction, review and synthesis",
                "responses.parse with Pydantic text_format; model configurable; store=False",
            ),
            (
                "LangGraph",
                "Stateful orchestration, Send, subgraphs and interrupts",
                "StateGraph plus SqliteSaver checkpoints",
            ),
            (
                "Streamlit",
                "Local research interface and periodic progress display",
                "Background service polled with st.fragment",
            ),
            (
                "SQLite",
                "Checkpoints, durable run state, budgets, cache and events",
                "Atomic reservation before every external attempt",
            ),
            (
                "Pydantic / HTTPX / python-dotenv",
                "Schema validation, HTTP requests and configuration",
                "No credentials in generated reports or event diagnostics",
            ),
        ],
    )
    para(
        doc,
        "The configured default analysis model is gpt-4.1-mini. A replacement must support the Responses API structured-output contract and be available to the account. Search and model calls may incur provider charges; this overview does not estimate current prices.",
    )
    para(
        doc,
        "Official-domain allowlists guide initial pricing, features and positioning searches. News is searched broadly with an explicit date range. Geography informs queries and overlap scoring; it does not guarantee that every result is geographically local. There is no automatic alternate search provider, vector database, or employee-policy RAG dependency in this application.",
    )
    code(
        doc,
        "YOU_API_KEY=<private You.com key>\nOPENAI_API_KEY=<private OpenAI key>\nOPENAI_MODEL=gpt-4.1-mini\nOPENAI_MAX_ATTEMPTS=5\nOPENAI_MAX_CONCURRENCY=1\nRESEARCH_DATA_DIR=artifacts",
    )
    para(
        doc,
        ".env is excluded from Git. Agentic_AI_Project/.vscode/settings.json enables python.terminal.useEnvFile and points python.envFile to ${workspaceFolder}/.env. New Python terminals apply that setting; restarting Streamlit is required to refresh cached service configuration.",
    )

    heading(doc, "5. Confidence, completeness and stop conditions")
    para(
        doc,
        "Competitive overlap and evidence confidence answer different questions. Ranking weights are product overlap 50%, customer overlap 35% and geography overlap 15%. Evidence confidence asks whether the selected company or extracted claim is actually supported by the available material.",
    )
    table(
        doc,
        ["Claim score component", "Weight", "Meaning"],
        [
            ("Source authority", "25%", "Official source versus external material"),
            ("Reviewer directness", "35%", "How directly the cited evidence supports the claim"),
            ("Content depth", "15%", "Full page, highlights or snippet"),
            ("Corroboration", "10%", "Distinct publishers and differing source text"),
            (
                "Freshness",
                "15%",
                "Current retrieval snapshot; dated in-window evidence required for news",
            ),
        ],
    )
    bullets(
        doc,
        [
            "Accept: valid citations, reviewer support, score ≥ 0.80 and no material conflict.",
            "Follow up: required categories are missing or claims remain below the acceptance gate.",
            "Escalate: identity/selection ambiguity, unresolved material conflicts, or supported profile claims averaging below 0.60 after bounded follow-up.",
            "Stop: quality/coverage met, no new evidence, maximum rounds, cancellation, deadline or exhausted call budget.",
        ],
    )
    para(
        doc,
        "Completeness is the fraction of four categories with accepted findings: pricing, features, positioning and news. It does not measure every feature or pricing tier. Profile confidence averages supported claim scores, so a high confidence score can coexist with missing categories. Scores are heuristics, not calibrated probabilities of truth.",
    )
    para(
        doc,
        "Conflicts cap confidence at 0.55; snippets alone cap it at 0.69. Invented evidence IDs, duplicate claim/assessment IDs, missing assessments and undated or out-of-window news cannot pass acceptance. Only accepted claims and a verified target baseline enter synthesis. An insight with any invalid claim reference is omitted.",
    )
    table(
        doc,
        ["Control", "Current default", "Scope"],
        [
            ("Search attempts", "40", "Whole run; retries counted"),
            ("Analysis attempts", "35", "Whole run; all agents and retries counted"),
            ("Active execution window", "600 seconds", "Run execution; renewed on resume"),
            (
                "OpenAI attempts per logical call",
                "5; configurable 1–10",
                "Initial attempt included; does not raise the run budget",
            ),
            (
                "OpenAI concurrent calls",
                "1; configurable 1–8",
                "Shared across branches and runs in the process",
            ),
            ("You.com attempts per call", "3", "Transient failures only"),
            (
                "Competitor research rounds",
                "Initial + up to 2 follow-ups",
                "Stops earlier if no improvement",
            ),
        ],
    )
    para(
        doc,
        "Retry-After seconds/HTTP dates and OpenAI retry-after-ms are honored with jitter without shortening the requested delay. Waits check cancellation/deadlines. A shared model slot and cooldown prevent queued model calls from bypassing a rate-limit delay, while web research remains parallel. Quota, billing, spend-limit, authentication and invalid-request failures require action and stop without retries.",
    )

    heading(doc, "6. Human review, checkpoint recovery and tracing")
    table(
        doc,
        ["Action", "Human input", "Effect"],
        [
            (
                "accept_partial",
                "Accept current caveats",
                "Continue while preserving gaps; disputed facts remain unverified",
            ),
            (
                "clarify",
                "Target product, market or buyer scope",
                "Repeat discovery with clarification",
            ),
            (
                "replace",
                "Three distinct competitor names/domains",
                "Research human-selected competitors with explicit labeling",
            ),
            (
                "retry",
                "Targeted direction during evidence review",
                "Rerun pending branches with remaining budgets",
            ),
            ("cancel", "Decision to stop", "End the research run"),
        ],
    )
    para(
        doc,
        "Selection review offers accept_partial, clarify, replace and cancel. Evidence review offers accept_partial, retry and cancel. A human decision resolves scope or authorizes incomplete delivery; it is not an automatic verification of the market claims.",
    )
    para(
        doc,
        "Retry saved run resumes the persisted LangGraph state. Completed nodes remain saved; successful search/model calls are cached within that run and reused when their inputs match. Unfinished nodes may restart, but matching cached calls are not paid requests again. Failed calls are not cached. Remaining overall budgets are retained, while the active time window is renewed. A finished partial run cannot currently be resumed with a larger budget; start a new run for that.",
    )
    para(
        doc,
        "checkpoints.sqlite stores graph state. research.sqlite stores run metadata, results, atomic budgets, caches and collaboration events. The trace records delegation, query planning, tool starts/completions, evidence handoffs, quality scores, queued model slots, retries, human decisions and stop reasons. OpenAI diagnostics include HTTP status, machine error code/type, exception type and request ID when available. Raw provider prose, headers and credentials are excluded.",
    )
    para(
        doc,
        "A crash after an external response but before saving its result can repeat that request on recovery. Checkpointing provides resumability, not exactly-once provider billing. Only one Streamlit process should operate on a data directory at a time.",
    )

    doc.add_page_break()
    heading(doc, "7. Prompts used during vibe coding")
    para(
        doc,
        "This section records the visible development instructions provided to the coding assistant. The quoted excerpts come from our conversation. They are different from the application agent prompts in Section 8; private assistant reasoning is not part of this record.",
    )
    heading(doc, "Initial project request — key verbatim requirements", 2)
    quote(
        doc,
        "Design a multi-agent market research system for competitor analysis:\nAgent 1 discovers a company’s top 3 competitors using the you.com Search API\nAgent 2 gathers fresh web and news data for each competitor\nAgent 3 extracts pricing, core features, market positioning, and recent news\nand an Orchestrator Agent coordinates the full research pipeline and compiles the findings into structured competitor analysis briefings and displays the analysis on the UI.",
    )
    quote(
        doc,
        "Add other subagents as needed. Modeled on real competitive intelligence workflows used by product teams, strategy teams, consultants, and founders. This project demonstrates autonomous web research, tool use, delegation, stateful orchestration, and structured analysis. Build with LangGraph.",
    )
    quote(
        doc,
        "Use you.com as the primary search source and present the results through a simple interface such as Streamlit.",
    )
    quote(
        doc,
        "In the Observation & Feedback Loop while running agents, take care of retries, stop condition, or escalate to a human if necessary based on relevant confidence scores.",
    )
    quote(doc, "Have a way of tracing how the agents worked together to solve the problem.")
    quote(
        doc,
        "Do not build directly. First, brainstorm with me, tell me your plan and only start building once I give you the go-ahead",
    )
    para(
        doc,
        "Result: the assistant first reviewed project context and official documentation, proposed the agent responsibilities and quality controls, and waited for approval before implementing.",
    )
    heading(doc, "Scope approval", 2)
    quote(
        doc,
        "Yes, go ahead with the focused Saas version, but having capabilities to extend further to other industries.\nYes, it should live in Agentic_AI_Project",
    )
    para(
        doc,
        "Result: a separate application was built in the initially empty Agentic_AI_Project workspace, with an industry profile registry and an offline demo.",
    )
    heading(doc, "Configuration and debugging prompts", 2)
    table(
        doc,
        ["Development prompt / excerpt", "Result"],
        [
            (
                '"Show me how to add both API keys in .env"',
                "Documented local key setup and app restart",
            ),
            (
                '"Enable python.terminal.useEnvFile to use environment variables from .env files in terminals"',
                "Added workspace Python environment-file settings",
            ),
            (
                '"Got this error, are the attempts being limited somewhere?"',
                "Distinguished per-call retries from overall analysis budget and inspected failed-run metadata",
            ),
            (
                '"How should I handle this error"',
                "Recommended classified errors, safe diagnostics, longer server-directed waits and limited model concurrency",
            ),
            (
                '"Can you go ahead and put in the recommended changes?"',
                "Implemented retry configuration, shared model gate/cooldown, terminal quota handling and UI failure details",
            ),
            (
                '"Will Retry Saved Run use researched data ... due to checkpointing?" [excerpt]',
                "Explained saved-node/cache reuse and remaining-budget behavior",
            ),
            (
                '"What is the Human in the Loop doing here?"',
                "Explained selection-review actions and uncertainty preservation",
            ),
            (
                '"I got Model call budget exhausted while testing for Shopify. Where is the limit being applied?"',
                "Confirmed the saved 16-call run budget and atomic reservation mechanism",
            ),
        ],
    )
    para(
        doc,
        "The iterations followed a reviewable cycle: clarify scope → build → run deterministic checks → inspect observed failures → make a bounded change → validate again. Configuration and recovery questions exposed product explanations that needed to accompany the technical implementation.",
    )

    doc.add_page_break()
    heading(doc, "8. Runtime prompts used by the agents")
    para(
        doc,
        "The following prompts are extracted from current source code, rather than reconstructed from memory. Each live model call receives the common system instruction plus its stage instruction and a JSON payload. The extractor prompt below includes the current SaaS pricing-rule suffix. Environment keys are never included in those research payloads.",
    )
    heading(doc, "Shared evidence-grounding instruction", 2)
    quote(doc, INVENTORY["system"])
    para(
        doc,
        "Source: market_research/providers.py, SYSTEM. Purpose: treat web material as untrusted evidence, prevent recalled facts from becoming findings, and require supported citations and explicit uncertainty.",
    )
    order = [
        "ResearchPlan",
        "Discovery",
        "DiscoveryReview",
        "CollectionPlan",
        "Extraction",
        "Review",
        "Briefing",
    ]
    prompts = sorted(INVENTORY["prompts"], key=lambda p: order.index(p["schema"]))
    explanations = {
        "ResearchPlan": "Payload: user research request and industry rules. Output: objective, discovery queries and comparison focus.",
        "Discovery": "Payload: request, discovery evidence and optional human clarification. Output: target baseline, candidate overlap scores and selection citations.",
        "DiscoveryReview": "Payload: proposed discovery result and all discovery evidence. Output: independent support judgments for the target and candidate identities.",
        "CollectionPlan": "Payload: competitor identity, scope and industry profile. Output: category-specific search tasks; deterministic guards ensure required categories remain present.",
        "Extraction": "Payload: competitor, scope, taxonomy and collected source records. Output: cited claims with nullable price/date fields plus reported gaps.",
        "Review": "Payload: extracted claims, competitor identity, scope and all source records. Output: support/directness/conflict assessments and targeted follow-up queries.",
        "Briefing": "Payload: accepted competitor claims and verified target baseline. Output: strategic insights referring to domain:claim_id identifiers.",
    }
    for i, prompt in enumerate(prompts, 1):
        heading(doc, f"8.{i} {prompt['agent']} — {prompt['schema']}", 2)
        quote(doc, prompt["prompt"])
        para(doc, explanations[prompt["schema"]])
        para(
            doc,
            f"Source: market_research/graph.py:{prompt['line']}. Prompt wording can guide the model; local schema and graph checks enforce operational boundaries.",
        )

    heading(doc, "9. Iterations and implementation changes")
    para(
        doc,
        "The following is the observed development sequence. It does not claim unrecorded experiments, model comparisons or quantitative accuracy benchmarks.",
    )
    table(
        doc,
        ["Iteration", "What was tried or changed", "Observed outcome"],
        [
            (
                "Design before implementation",
                "Reviewed existing Hybrid RAG workspace and official LangGraph/You.com documentation; proposed evidence-first workflow",
                "User approved a separate SaaS application in Agentic_AI_Project",
            ),
            (
                "First graph + UI",
                "Built schemas, provider adapters, SQLite budgets/cache, parallel branches, export rendering and demo",
                "Initial end-to-end suite: 19 passing tests",
            ),
            (
                "Quality + UI hardening",
                "Added Streamlit flow tests, duplicate/date checks, cache timestamp preservation and selection-caveat handling",
                "Expanded suite: 26 passing tests",
            ),
            (
                "Autonomous planning + independent selection review",
                "Added structured orchestrator/collector plans and reviewer validation of target/candidate identities; added real-SDK mocked HTTP checks",
                "Expanded suite: 31 passing tests; normal demo used 14 searches and 13 analysis calls",
            ),
            (
                "API-key ergonomics",
                "Documented .env setup and enabled VS Code Python terminal environment-file use",
                "Local setup instructions clarified; credentials remained outside reports/Git",
            ),
            (
                "Observed retry failure",
                "Inspected a failed live run and explained three attempts per call versus the shared run budget",
                "Original retry trace lacked the failure's HTTP/error classification",
            ),
            (
                "Retry and concurrency improvement",
                "Added classified diagnostics, configurable five-attempt default, one shared model slot, server-directed delays and terminal quota handling",
                "45 passing tests; application restarted with saved checkpoints retained",
            ),
            (
                "Shopify budget investigation",
                "Inspected saved request limits and model-call counter",
                "Run exhausted 16 of 16 configured analysis attempts and returned partial results",
            ),
        ],
    )
    para(
        doc,
        "The original three-attempt limit was not itself proof of an OpenAI outage. The old trace recorded delays but not status/error codes, so the exact historical cause remains unknown. The improved implementation preserves those details for future failures; it does not retroactively diagnose old events.",
    )

    heading(doc, "10. Observations from recorded runs")
    rows = []
    chosen = []
    for name in ["Shopify", "Workday", "Salesforce", "FlowPilot"]:
        found = next((r for r in INVENTORY["runs"] if r["company"] == name), None)
        if found:
            chosen.append(found)
            rows.append(
                (
                    name,
                    "Fictional demo" if found["demo"] else "Live mode",
                    found["status"],
                    f"{found['queries']} / {found['max_queries']}",
                    f"{found['calls']} / {found['max_calls']}",
                )
            )
    table(doc, ["Company", "Mode", "Saved status", "Searches / limit", "Model calls / limit"], rows)
    para(
        doc,
        "These counts come from the local run database when this document was generated. They describe execution and configured limits, not search quality, competitor correctness or the truth of extracted claims. Subsequent retries or new runs can change the saved metadata.",
    )
    bullets(
        doc,
        [
            "Shopify: the 16-call limit was lower than the application's default 35, demonstrating why saved configuration must be inspected before changing retry code.",
            "Salesforce: the trace includes the earlier generic OpenAI retry-exhaustion message; the current saved status is partial, showing that later recovery can preserve usable findings without declaring the briefing complete.",
            "Workday: a complete saved workflow used 13 model calls within its 15-call limit. That leaves little retry/follow-up headroom and is not a recommended universal budget.",
            "FlowPilot: the fictional demo used 14 searches and 13 model calls in the expanded workflow. The earliest saved demo used eight model calls before additional planning/review stages were introduced.",
        ],
    )

    heading(doc, "11. Learnings and practical implications")
    table(
        doc,
        ["Learning", "Workflow implication"],
        [
            (
                "Retries and budgets are different controls",
                "Per-call resilience cannot override run-wide cost/attempt limits; display both clearly",
            ),
            (
                "Parallel agents share provider capacity",
                "Keep source gathering parallel, but serialize model requests initially and honor shared cooldowns",
            ),
            (
                "A source ID is not evidence support",
                "Independent review checks relevance/identity; deterministic checks reject missing or invented IDs",
            ),
            (
                "Confidence is different from coverage",
                "Show missing categories even when accepted claims have high confidence",
            ),
            (
                "Pricing needs a comparison basis",
                "Preserve currency, seats/usage, region and billing commitment; avoid false headline-price comparisons",
            ),
            (
                "Fresh search does not prove fresh news",
                "Check source publication dates and event dates; disclose unknown dates",
            ),
            (
                "Human review should resolve a concrete issue",
                "Provide scope clarification, replacement choices or targeted retry; preserve caveats after acceptance",
            ),
            (
                "Durability must include intermediate tool results",
                "Checkpoint graph state and cache successful calls to make recovery efficient",
            ),
            (
                "Operational diagnostics are part of the product",
                "Actionable errors and inspectable traces are more useful than a generic unavailable message",
            ),
            (
                "Prompting needs local enforcement",
                "Use structured outputs, reducers, budget reservations and route rules in addition to natural-language instructions",
            ),
        ],
    )
    heading(doc, "Limitations and extension opportunities", 2)
    bullets(
        doc,
        [
            "Model-assisted grounding and selection still require representative live evaluation; passing tests are not an accuracy benchmark.",
            "Exact duplicate/syndicated text is deduplicated; near-duplicates and subtle source conflicts can remain.",
            "Synthesis filters reference validity, but that alone cannot prove every strategic interpretation is substantively justified.",
            "The prototype is local and shares saved research within a data directory; it has no per-user authentication or access controls.",
            "Reports are snapshots. Continuous competitor monitoring, alerting and historical diffs would require additional scheduling and state design.",
            "Other industries can reuse profiles and orchestration; new analysis categories need coordinated schema, collector, quality and rendering changes.",
            "A future UX improvement could let a user explicitly extend a finished partial run's budget while recording that authorization; the current implementation requires a new run.",
        ],
    )

    heading(doc, "12. Validation and reproducibility")
    para(
        doc,
        "The latest recorded application validation completed with 45 passing tests, plus successful lint and formatting checks. Tests cover the actual LangGraph pipeline, normal/conflict/ambiguity flows, checkpoint recovery, concurrent run isolation, atomic budgets, fabricated citations, dates, cache timestamps, the You.com contract, real OpenAI SDK parsing with mocked HTTP, long Retry-After waits, quota handling, shared cooldowns, queued cancellation and Streamlit review flows.",
    )
    para(
        doc,
        "Those checks made no paid API calls. Saved live-mode runs provide operational evidence but were not manually benchmarked here against a labeled competitor-analysis dataset. The Word document's prompt inventory is extracted from source code; architecture diagrams are generated programmatically and embedded as images.",
    )
    code(
        doc,
        "cd /Users/sskamble/VSCode_Projects/Agentic_AI_Project\nsource .venv/bin/activate\nstreamlit run app.py\n\npython -m pytest -q\nruff check .\nruff format --check .\npython research.py --output artifacts/demo-report",
    )
    para(
        doc, "Document regeneration (uses existing libraries across the two project environments):"
    )
    code(
        doc,
        ".venv/bin/python -m reports.build_document_assets\n../Hybrid_RAG_Project/.venv/bin/python reports/build_project_document.py",
    )
    heading(doc, "13. References and source map")
    para(
        doc,
        "Project references: README.md; requirements.txt; requirements-lock.txt; app.py; research.py; market_research/graph.py, models.py, providers.py, quality.py, storage.py, service.py, industries.py, demo.py and rendering.py; tests/; artifacts/research.sqlite; reports/document_inventory.json. The report references the current implementation, not a deployment or release commit.",
    )
    for title, url in [
        ("You.com Search API reference", "https://you.com/docs/api-reference/search/v1-search"),
        ("You.com search and extraction guide", "https://you.com/docs/guides/search"),
        ("LangGraph Graph API", "https://docs.langchain.com/oss/python/langgraph/graph-api"),
        ("LangGraph persistence", "https://docs.langchain.com/oss/python/langgraph/persistence"),
        ("LangGraph interrupts", "https://docs.langchain.com/oss/python/langgraph/interrupts"),
        (
            "OpenAI structured outputs",
            "https://developers.openai.com/api/docs/guides/structured-outputs",
        ),
        ("OpenAI rate limits", "https://developers.openai.com/api/docs/guides/rate-limits"),
        ("OpenAI error codes", "https://developers.openai.com/api/docs/guides/error-codes"),
        (
            "Streamlit fragment API",
            "https://docs.streamlit.io/develop/api-reference/execution-flow/st.fragment",
        ),
    ]:
        link(doc, title, url)
    doc.save(DOCUMENT)
    check = Document(DOCUMENT)
    assert len(check.inline_shapes) == 2
    assert len(check.tables) >= 8
    assert all(
        any(title in p.text for p in check.paragraphs)
        for title in [
            "Project overview",
            "APIs, libraries",
            "Prompts used during vibe coding",
            "Runtime prompts",
            "Iterations",
            "Learnings",
            "Validation",
        ]
    )
    print(
        f"Created {DOCUMENT.name}: {len(check.paragraphs)} paragraphs, {len(check.tables)} tables, "
        f"{len(check.inline_shapes)} diagrams, {DOCUMENT.stat().st_size:,} bytes."
    )


if __name__ == "__main__":
    build()
